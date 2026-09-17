"""
Builds the downloadable Excel template for each table.

The template has:
  * "Data"  - the sheet users fill in, with the exact spec headers and
              in-cell dropdowns for Project / Implementer / Intervention /
              District / Sector / Cell / Village / Gender / Youth / and any
              fixed-choice fields declared in schemas.py (ftype "choice").
  * "Lists" - the option lists that back those dropdowns.
  * "Location_Lookup" - full District > Sector > Cell > Village reference.
  * "Intervention_Lookup" - full Implementer > Intervention reference.

Implementer -> Intervention is a cascading dropdown. District -> Sector ->
Cell -> Village is also a cascading dropdown using named ranges + INDIRECT,
so each location level is filtered by the value selected in the previous
level on the same row.
"""
from __future__ import annotations
import io
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment

import config
import schemas
import reference_data as ref

MAX_ROWS = 2000                       # rows the validation applies to
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
REQ_FILL = PatternFill("solid", fgColor="C00000")


def _safe_name(s: str) -> str:
    """Excel defined-name safe token (letters, digits, underscore; not leading digit)."""
    out = re.sub(r"[^A-Za-z0-9_]", "_", str(s))
    if not out or out[0].isdigit():
        out = "_" + out
    return out[:250]


def _write_list_col(ws, col_idx: int, title: str, values: list[str]) -> str:
    """Write a vertical list, return its absolute range like 'Lists'!$A$2:$A$9."""
    letter = get_column_letter(col_idx)
    ws.cell(row=1, column=col_idx, value=title).font = Font(bold=True)
    for i, v in enumerate(values, start=2):
        ws.cell(row=i, column=col_idx, value=v)
    last = max(len(values) + 1, 2)
    return f"'{ws.title}'!${letter}$2:${letter}${last}"


def _write_named_location_list(
    wb: Workbook,
    ws,
    col_idx: int,
    title: str,
    values: list[str],
    name: str,
) -> str | None:
    """Write a location list and register it as an Excel defined name."""
    values = sorted({str(v).strip() for v in values if pd.notna(v) and str(v).strip()})
    if not values:
        return None

    rng = _write_list_col(ws, col_idx, title, values)
    wb.defined_names.add(DefinedName(_safe_name(name), attr_text=rng))
    return rng


def build_template(table_key: str) -> bytes:
    cfg = schemas.TABLES[table_key]
    fields = cfg["fields"]
    labels = [fl["label"] for fl in fields]
    ftypes = [fl["ftype"] for fl in fields]

    wb = Workbook()

    # ---------------- Data sheet ----------------
    ws = wb.active
    ws.title = "Data"
    for j, fl in enumerate(fields, start=1):
        c = ws.cell(row=1, column=j, value=fl["label"])
        c.font = HDR_FONT
        c.fill = REQ_FILL if fl["req"] else HDR_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = max(14, min(34, len(fl["label"]) + 4))
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 34

    # ---------------- Lists sheet ----------------
    lists = wb.create_sheet("Lists")
    col = 1
    rng_project = _write_list_col(lists, col, "Project", config.PROJECTS); col += 1
    implementers = ref.implementers()
    rng_impl = _write_list_col(lists, col, "Implementer", implementers); col += 1
    rng_gender = _write_list_col(lists, col, "Gender", config.GENDER_OPTIONS); col += 1
    rng_youth = _write_list_col(lists, col, "Youth", config.YOUTH_OPTIONS); col += 1

    # Fixed-choice fields declared per-table in schemas.py (ftype "choice"),
    # e.g. Beneficiary Category, Age_Group, Employment Contract, etc.
    # Each gets its own flat dropdown list, keyed by column name.
    choice_ranges: dict[str, str] = {}
    for fl in fields:
        if fl["ftype"] == "choice" and fl.get("options"):
            rng = _write_list_col(lists, col, fl["label"], fl["options"]); col += 1
            choice_ranges[fl["col"]] = rng

    # Location hierarchy.  These named ranges are used by the cascading
    # District -> Sector -> Cell -> Village validations below.
    location_df = ref.load_locations().copy()
    location_df.columns = [str(c).strip() for c in location_df.columns]
    location_cols = ["District", "Sector", "Cell", "Village"]

    for lc in location_cols:
        if lc not in location_df.columns:
            raise ValueError(
                f"Location lookup is missing required column: {lc}. "
                f"Found: {list(location_df.columns)}"
            )
        location_df[lc] = location_df[lc].fillna("").astype(str).str.strip()

    location_df = location_df[
        location_df["District"].ne("")
    ].drop_duplicates()

    rng_district = _write_named_location_list(
        wb, lists, col, "District", location_df["District"].tolist(), "District"
    )
    col += 1

    # One named range for each District -> Sector, District/Sector -> Cell,
    # and District/Sector/Cell -> Village combination.
    location_ranges: dict[tuple[str, ...], str] = {}

    for district, grp_d in location_df.groupby("District", sort=True):
        district_name = _safe_name(district)
        sectors = sorted({v for v in grp_d["Sector"] if v})
        if sectors:
            name = f"SEC_{district_name}"
            rng = _write_named_location_list(
                wb, lists, col, f"SEC_{district}", sectors, name
            )
            col += 1
            if rng:
                location_ranges[(district,)] = name

        for sector, grp_s in grp_d.groupby("Sector", sort=True):
            if not sector:
                continue

            sector_name = _safe_name(sector)
            cells = sorted({v for v in grp_s["Cell"] if v})
            if cells:
                name = f"CEL_{district_name}_{sector_name}"
                rng = _write_named_location_list(
                    wb, lists, col, f"CEL_{district}_{sector}", cells, name
                )
                col += 1
                if rng:
                    location_ranges[(district, sector)] = name

            for cell, grp_c in grp_s.groupby("Cell", sort=True):
                if not cell:
                    continue

                cell_name = _safe_name(cell)
                villages = sorted({v for v in grp_c["Village"] if v})
                if villages:
                    name = f"VIL_{district_name}_{sector_name}_{cell_name}"
                    rng = _write_named_location_list(
                        wb,
                        lists,
                        col,
                        f"VIL_{district}_{sector}_{cell}",
                        villages,
                        name,
                    )
                    col += 1
                    if rng:
                        location_ranges[(district, sector, cell)] = name

    tree_df = ref.load_trees()
    rng_origin = rng_ttype = None
    tree_species_ranges: dict[str, str] = {}
    if tree_df is not None:
        rng_origin = _write_list_col(lists, col, "Tree Orgine", ref.tree_origins()); col += 1
        rng_ttype = _write_list_col(lists, col, "Tree Type",
                                    sorted({v for v in tree_df["Tree Type"] if v})); col += 1
        # one named range per (origin, type) combo, for a true 2-level species cascade
        for (origin, ttype), grp in tree_df.groupby(["Tree Origin", "Tree Type"]):
            species = sorted({v for v in grp["Tree species"] if v})
            if not species:
                continue
            rng = _write_list_col(lists, col, f"SP_{origin}_{ttype}", species)
            combo_name = _safe_name(f"{origin}_{ttype}")
            wb.defined_names.add(DefinedName(combo_name, attr_text=rng))
            tree_species_ranges[combo_name] = rng
            col += 1

    # per-implementer intervention lists (for the INDIRECT cascade)
    iv = ref.load_interventions()
    for impl in implementers:
        opts = sorted({v for v in iv.loc[iv["Implementer"] == impl, "Intervention"] if v})
        if not opts:
            continue
        rng = _write_list_col(lists, col, f"IV_{impl}", opts)
        wb.defined_names.add(DefinedName(_safe_name(impl), attr_text=rng))
        col += 1

    # ---------------- Lookup sheets ----------------
    loc_df = location_df
    ws_loc = wb.create_sheet("Location_Lookup")
    ws_loc.append(["District", "Sector", "Cell", "Village"])
    for c in ws_loc[1]:
        c.font = Font(bold=True)
    for row in loc_df.itertuples(index=False):
        ws_loc.append(list(row))
    for i, w in enumerate([18, 18, 20, 22], start=1):
        ws_loc.column_dimensions[get_column_letter(i)].width = w
    ws_loc.freeze_panes = "A2"
    ws_loc.auto_filter.ref = ws_loc.dimensions

    ws_iv = wb.create_sheet("Intervention_Lookup")
    ws_iv.append(["Implementer", "Intervention"])
    for c in ws_iv[1]:
        c.font = Font(bold=True)
    for row in iv.itertuples(index=False):
        ws_iv.append(list(row))
    ws_iv.column_dimensions["A"].width = 20
    ws_iv.column_dimensions["B"].width = 60
    ws_iv.freeze_panes = "A2"
    ws_iv.auto_filter.ref = ws_iv.dimensions

    if tree_df is not None:
        ws_tr = wb.create_sheet("Tree_Lookup")
        ws_tr.append(["Tree Orgine", "Tree Type", "Tree species"])
        for c in ws_tr[1]:
            c.font = Font(bold=True)
        for row in tree_df.itertuples(index=False):
            ws_tr.append(list(row))
        for i, w in enumerate([18, 18, 28], start=1):
            ws_tr.column_dimensions[get_column_letter(i)].width = w
        ws_tr.freeze_panes = "A2"

    # ---------------- Validations on Data sheet ----------------
    def add_dv(col_idx: int, formula: str, prompt: str = ""):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
        dv.error = "Pick a value from the dropdown list."
        dv.errorTitle = "Invalid entry"
        if prompt:
            dv.prompt = prompt
            dv.promptTitle = "Reference list"
        ws.add_data_validation(dv)
        letter = get_column_letter(col_idx)
        dv.add(f"{letter}2:{letter}{MAX_ROWS}")

    impl_col_idx = None
    district_col_idx = None
    sector_col_idx = None
    cell_col_idx = None
    origin_col_idx = None
    ttype_col_idx = None
    for j, fl in enumerate(fields, start=1):
        ft = fl["ftype"]
        if ft == "project":
            add_dv(j, rng_project)
        elif ft == "implementer":
            impl_col_idx = j
            add_dv(j, rng_impl)
        elif ft == "gender":
            add_dv(j, rng_gender)
        elif ft == "youth":
            add_dv(j, rng_youth)
        elif ft == "choice" and fl["col"] in choice_ranges:
            add_dv(j, choice_ranges[fl["col"]])
        elif ft == "district":
            district_col_idx = j
            add_dv(j, rng_district, "Select a District first. Sector, Cell and Village will filter from it.")
        elif ft == "sector":
            sector_col_idx = j
        elif ft == "cell":
            cell_col_idx = j
        elif ft == "village":
            # Validation is added below after all location columns are known.
            pass
        elif ft == "tree_origin" and rng_origin:
            origin_col_idx = j
            add_dv(j, rng_origin)
        elif ft == "tree_type" and rng_ttype:
            ttype_col_idx = j
            add_dv(j, rng_ttype)

    # ------------------------------------------------------------------
    # Cascading location validation:
    # District -> Sector -> Cell -> Village
    #
    # Each dependent dropdown uses the selected values from the same row.
    # Named ranges are generated above from Location_Lookup.
    # ------------------------------------------------------------------
    if district_col_idx:
        district_letter = get_column_letter(district_col_idx)

        for j, ft in enumerate(ftypes, start=1):
            if ft == "sector":
                formula = f'=INDIRECT("SEC_"&SUBSTITUTE(${district_letter}2," ","_"))'
                dv = DataValidation(
                    type="list",
                    formula1=formula,
                    allow_blank=True,
                    showDropDown=False,
                )
                dv.error = "Choose a Sector that belongs to the selected District."
                dv.errorTitle = "Invalid District/Sector"
                dv.prompt = "Select a District first — this list is filtered to that District."
                dv.promptTitle = "Depends on District"
                ws.add_data_validation(dv)
                letter = get_column_letter(j)
                dv.add(f"{letter}2:{letter}{MAX_ROWS}")

    if district_col_idx and sector_col_idx:
        district_letter = get_column_letter(district_col_idx)
        sector_letter = get_column_letter(sector_col_idx)

        for j, ft in enumerate(ftypes, start=1):
            if ft == "cell":
                formula = (
                    f'=INDIRECT("CEL_"'
                    f'&SUBSTITUTE(${district_letter}2," ","_")'
                    f'&"_"&SUBSTITUTE(${sector_letter}2," ","_"))'
                )
                dv = DataValidation(
                    type="list",
                    formula1=formula,
                    allow_blank=True,
                    showDropDown=False,
                )
                dv.error = "Choose a Cell that belongs to the selected District and Sector."
                dv.errorTitle = "Invalid District/Sector/Cell"
                dv.prompt = "Select District and Sector first — this list is filtered to that location."
                dv.promptTitle = "Depends on District + Sector"
                ws.add_data_validation(dv)
                letter = get_column_letter(j)
                dv.add(f"{letter}2:{letter}{MAX_ROWS}")

    if district_col_idx and sector_col_idx and cell_col_idx:
        district_letter = get_column_letter(district_col_idx)
        sector_letter = get_column_letter(sector_col_idx)
        cell_letter = get_column_letter(cell_col_idx)

        for j, ft in enumerate(ftypes, start=1):
            if ft == "village":
                formula = (
                    f'=INDIRECT("VIL_"'
                    f'&SUBSTITUTE(${district_letter}2," ","_")'
                    f'&"_"&SUBSTITUTE(${sector_letter}2," ","_")'
                    f'&"_"&SUBSTITUTE(${cell_letter}2," ","_"))'
                )
                dv = DataValidation(
                    type="list",
                    formula1=formula,
                    allow_blank=True,
                    showDropDown=False,
                )
                dv.error = "Choose a Village that belongs to the selected District, Sector and Cell."
                dv.errorTitle = "Invalid District/Sector/Cell/Village"
                dv.prompt = "Select District, Sector and Cell first — this list is filtered to that location."
                dv.promptTitle = "Depends on District + Sector + Cell"
                ws.add_data_validation(dv)
                letter = get_column_letter(j)
                dv.add(f"{letter}2:{letter}{MAX_ROWS}")

    # cascading Intervention: depends on the Implementer cell in the same row
    if impl_col_idx:
        for j, ft in enumerate(ftypes, start=1):
            if ft == "intervention":
                impl_letter = get_column_letter(impl_col_idx)
                formula = f'=INDIRECT(SUBSTITUTE(${impl_letter}2," ","_"))'
                dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
                dv.error = "Choose an Intervention that belongs to the Implementer in this row."
                dv.errorTitle = "Invalid Implementer/Intervention pair"
                dv.prompt = "Fill the Implementer cell first — this list filters to that implementer."
                dv.promptTitle = "Depends on Implementer"
                ws.add_data_validation(dv)
                letter = get_column_letter(j)
                dv.add(f"{letter}2:{letter}{MAX_ROWS}")

    # cascading Tree species: depends on BOTH the Tree Origin and Tree Type
    # cells in the same row (two-level cascade via a combo-keyed named range)
    if tree_species_ranges and origin_col_idx and ttype_col_idx:
        origin_letter = get_column_letter(origin_col_idx)
        ttype_letter = get_column_letter(ttype_col_idx)
        for j, ft in enumerate(ftypes, start=1):
            if ft == "tree_species":
                formula = (
                    f'=INDIRECT(SUBSTITUTE(${origin_letter}2," ","_")'
                    f'&"_"&SUBSTITUTE(${ttype_letter}2," ","_"))'
                )
                dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
                dv.error = "Choose a species that matches the Tree Orgine/Category AND Tree Type in this row."
                dv.errorTitle = "Invalid tree combination"
                dv.prompt = "Fill Tree Orgine/Category and Tree Type first — this list filters to that combination."
                dv.promptTitle = "Depends on Tree Orgine + Tree Type"
                ws.add_data_validation(dv)
                letter = get_column_letter(j)
                dv.add(f"{letter}2:{letter}{MAX_ROWS}")

    # ---------------- Instructions ----------------
    ws_help = wb.create_sheet("READ_ME", 0)
    help_rows = [
        [f"{cfg['title']} — bulk upload template"],
        [],
        ["1.", "Fill in the 'Data' sheet only. One row per record. Do not rename or reorder the headers."],
        ["2.", "Red headers are REQUIRED. Blue headers are optional."],
        ["3.", "Project, Implementer, Gender, Youth, District, Sector, Cell, Village, Tree Orgine/Category, "
               "Tree Type, and any fixed-choice column (e.g. Beneficiary Category, Age_Group, "
               "Employment Contract) have in-cell dropdowns."],
        ["4.", "Intervention filters to the Implementer in the same row. Tree species filters to the Tree Orgine/Category AND Tree Type in the same row — fill those two first."],
        ["5.", "Location is cascading: select District first, then Sector, then Cell, then Village. Each dropdown is filtered to the previous selections in the same row."],
        ["6.", "Dates: use a real Excel date (YYYY-MM-DD). Numbers: digits only, no units or commas."],
        ["7.", "Leave blank for unknown values. Do not type 'N/A' or '-' in number or date columns."],
        ["8.", "Save as .xlsx and upload it on the table's 'Bulk upload' tab."],
        [],
        ["Note", "Rows that fail validation are never inserted. You get a list of them with the exact reason."],
    ]
    for r in help_rows:
        ws_help.append(r)
    ws_help["A1"].font = Font(bold=True, size=14)
    ws_help.column_dimensions["A"].width = 8
    ws_help.column_dimensions["B"].width = 105
    for row in ws_help.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
